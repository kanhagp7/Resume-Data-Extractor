from flask import Flask, render_template, request, send_file
import os
import re
import openpyxl
from PyPDF2 import PdfReader
from docx import Document

app = Flask(__name__)

def extract_information_from_cv(cv_path):
    _, file_extension = os.path.splitext(cv_path)
    
    if file_extension == '.pdf':
        with open(cv_path, 'rb') as f:
            pdf_reader = PdfReader(f)
            text = ''
            for page in pdf_reader.pages:
                text += page.extract_text()
    elif file_extension == '.docx':
        doc = Document(cv_path)
        text = ''
        for paragraph in doc.paragraphs:
            text += paragraph.text
    else:
        # Unsupported file type
        return [], [], ''

    email = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    phone = re.findall(r'\b\d{10}\b', text)
    return email, phone, text
def process_cvs(cvs):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Email", "Phone", "Text"])
    
    for cv in cvs:
        email, phone, text = extract_information_from_cv(cv)
        # Convert list of emails to a comma-separated string
        email_str = ', '.join(email)
        phone_str = ', '.join(phone)
        worksheet.append([email_str, phone_str, text])
    
    excel_filename = "cv_information.xlsx"
    workbook.save(excel_filename)
    return excel_filename


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload():
    uploaded_files = request.files.getlist("file")
    cv_paths = []
    for file in uploaded_files:
        filename = file.filename
        cv_path = os.path.join("uploads", filename)
        file.save(cv_path)
        cv_paths.append(cv_path)
    
    excel_file = process_cvs(cv_paths)
    return send_file(excel_file, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True)
